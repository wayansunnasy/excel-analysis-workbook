"""The proof in the README, as a script: build the example, recalculate it in LibreOffice, paste
50 rows under the table, recalculate again, and print both columns of figures.

Needs LibreOffice on the PATH as `soffice`. Run from anywhere:
    python tests/paste_proof.py
"""
import datetime
import pathlib
import subprocess
import sys
import tempfile

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import build_analysis_workbook as b
from openpyxl import load_workbook


def recalculate(path, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True, timeout=300)
    return load_workbook(outdir / path.name, data_only=True)["Analysis"]


def figures(ws):
    labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    grid = {l: ws.cell(i + 1, 2).value for i, l in enumerate(labels) if l}
    first = labels.index("By category") + 3
    other = labels.index("Other") + 1
    grid["Category block, all rows including Other"] = sum(ws.cell(r, 2).value for r in range(first, other + 1))
    grid["Category block, Other row"] = ws.cell(other, 2).value
    grid["Formula errors"] = sum(1 for row in ws.iter_rows() for c in row
                                 if isinstance(c.value, str) and c.value.startswith("#"))
    return grid


def main():
    work = pathlib.Path(tempfile.mkdtemp())
    built = work / "built.xlsx"
    b.build(HERE.parent / "examples" / "sample-products.csv", built)
    before = figures(recalculate(built, work / "a"))

    wb = load_workbook(built)
    data = wb["Data"]
    table = list(data.tables.values())[0]
    last = data.max_row
    for i in range(50):
        data.append([f"9990000{i:05d}", f"Pasted title {i}", "Pasted category", 100.00, "GBP", 4,
                     "yes", 1, "https://example.invalid", datetime.datetime(2026, 9, 23)])
    table.ref = f"A1:J{last + 50}"
    pasted = work / "pasted.xlsx"
    wb.save(pasted)
    after = figures(recalculate(pasted, work / "b"))

    keys = ["Rows of data", "Total price", "Total stock_count", "Latest scraped_at_utc",
            "Category block, Other row", "Category block, all rows including Other", "Formula errors"]
    print(f"{'Figure':44s} {'500 rows':>20s} {'after pasting 50':>20s}")
    for k in keys:
        print(f"{k:44s} {str(before.get(k)):>20s} {str(after.get(k)):>20s}")


if __name__ == "__main__":
    main()
